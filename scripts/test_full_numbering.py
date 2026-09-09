import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client, os

def apply_section_margins(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

def set_section_page_numbering(section, is_roman=False, start_num=None, suppress_first_page=True):
    apply_section_margins(section)
    section.different_first_page_header_footer = suppress_first_page
    
    sectPr = section._sectPr
    for child in list(sectPr):
        if child.tag.endswith('pgNumType'):
            sectPr.remove(child)
            
    fmt = 'lowerRoman' if is_roman else 'decimal'
    if start_num is not None:
        pg_xml = f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}" w:start="{start_num}"/>'
    else:
        pg_xml = f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}"/>'
    sectPr.append(parse_xml(pg_xml))
    
    # Configure Regular/Subsequent Pages Footer (aligned right)
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.text = ''
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p._p.append(fld)
    
    # Ensure First Page Footer is Empty if suppressed
    if suppress_first_page:
        fp_footer = section.first_page_footer
        fp_footer.is_linked_to_previous = False
        p_fp = fp_footer.paragraphs[0]
        p_fp.text = ''

doc = docx.Document()
# Section 1: Title & Bonafide
s1 = doc.sections[0]
apply_section_margins(s1)
s1.different_first_page_header_footer = True
s1.footer.is_linked_to_previous = False
s1.footer.paragraphs[0].text = ''
s1.first_page_footer.is_linked_to_previous = False
s1.first_page_footer.paragraphs[0].text = ''

doc.add_paragraph('Title Page')
doc.add_page_break()
doc.add_paragraph('Bonafide Certificate')
doc.add_page_break()
doc.add_paragraph('Design Constraints')

# Section 2: Acknowledgement & Front Matter (Roman starting at i)
s2 = doc.add_section()
set_section_page_numbering(s2, is_roman=True, start_num=1, suppress_first_page=False)
doc.add_paragraph('Acknowledgement (i)')
doc.add_page_break()
doc.add_paragraph('Abstract (ii)')
doc.add_page_break()
doc.add_paragraph('Table of Contents (iii)')

# Section 3: Chapter 1 (Arabic starting at 1, First page suppressed)
s3 = doc.add_section()
set_section_page_numbering(s3, is_roman=False, start_num=1, suppress_first_page=True)
doc.add_paragraph('Chapter 1 Page 1 (No number displayed)')
doc.add_page_break()
doc.add_paragraph('Chapter 1 Page 2 (Displays 2 on bottom right)')

# Section 4: Chapter 2 (Arabic continuous, First page suppressed)
s4 = doc.add_section()
set_section_page_numbering(s4, is_roman=False, start_num=None, suppress_first_page=True)
doc.add_paragraph('Chapter 2 Page 1 (No number displayed)')
doc.add_page_break()
doc.add_paragraph('Chapter 2 Page 2 (Displays 4 on bottom right)')

doc.save('test_full_numbering.docx')
print('Successfully saved test_full_numbering.docx')

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    wdoc = word.Documents.Open(os.path.abspath('test_full_numbering.docx'))
    print('Total pages in test doc:', wdoc.ComputeStatistics(2))
    print('Total sections:', wdoc.Sections.Count)
    for i in range(1, wdoc.Sections.Count + 1):
        sec = wdoc.Sections(i)
        print(f'Section {i}: DiffFirstPage={sec.PageSetup.DifferentFirstPageHeaderFooter}')
    wdoc.Close(False)
finally:
    word.Quit()
print("Verification complete!")
