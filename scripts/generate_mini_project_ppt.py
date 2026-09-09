import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

# --- Professional Dark Theme Palette ---
BG_DARK = RGBColor(15, 23, 42)        # Deep Navy / Slate 900
BG_CARD = RGBColor(30, 41, 59)        # Slate 800
BG_CARD_LIGHT = RGBColor(51, 65, 85)  # Slate 700
ACCENT_BLUE = RGBColor(56, 189, 248)  # Sky 400
ACCENT_CYAN = RGBColor(45, 212, 191)  # Teal 400
ACCENT_AMBER = RGBColor(251, 191, 36) # Amber 400
ACCENT_ROSE = RGBColor(244, 63, 94)   # Rose 500
ACCENT_GREEN = RGBColor(74, 222, 128) # Green 400
TEXT_MAIN = RGBColor(248, 250, 252)   # Slate 50
TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
BORDER_COLOR = RGBColor(71, 85, 105)  # Slate 600
WHITE = RGBColor(255, 255, 255)

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_blank_slide_with_bg():
        slide = prs.slides.add_slide(blank_layout)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return slide

    def add_header(slide, category, title, subtitle=None):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = 'Calibri'
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = 'Calibri'
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_MAIN

        if subtitle:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.25), Inches(11.7), Inches(0.35))
            tf_s = sub_box.text_frame
            tf_s.word_wrap = True
            tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle
            p_s.font.name = 'Calibri'
            p_s.font.size = Pt(12)
            p_s.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title=None, bg_color=BG_CARD, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.4))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = 'Calibri'
            p.font.size = Pt(13.5)
            p.font.bold = True
            p.font.color.rgb = ACCENT_BLUE
        return shape

    def add_bullet_list(slide, left, top, width, height, items, font_size=11.5, text_color=TEXT_MAIN):
        tb = slide.shapes.add_textbox(left, top, width, height)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(5)
            p.font.name = 'Calibri'
            p.font.size = Pt(font_size)
            
            if isinstance(item, tuple):
                prefix, text = item
                r1 = p.add_run()
                r1.text = prefix + ": "
                r1.font.bold = True
                r1.font.color.rgb = ACCENT_CYAN
                
                r2 = p.add_run()
                r2.text = text
                r2.font.bold = False
                r2.font.color.rgb = text_color
            else:
                r = p.add_run()
                r.text = item
                r.font.color.rgb = text_color

    def add_image_safe(slide, img_path, left, top, width=None, height=None):
        if os.path.exists(img_path):
            try:
                img = slide.shapes.add_picture(img_path, left, top, width=width, height=height)
                return img
            except Exception as e:
                print(f"Error adding image {img_path}: {e}")
        else:
            print(f"Image not found: {img_path}")
        return None

    # =========================================================================
    # SLIDE 1: TITLE & TEAM COVER SLIDE (WITH COLLEGE LOGO)
    # =========================================================================
    s1 = add_blank_slide_with_bg()
    
    # College Logo Container at Top Right
    logo_path = "docs/figures/dr_mgr_logo.png"
    if os.path.exists(logo_path):
        # Add white container pill for logo
        logo_bg = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.5), Inches(0.45), Inches(3.0), Inches(0.85))
        logo_bg.fill.solid()
        logo_bg.fill.fore_color.rgb = WHITE
        logo_bg.line.color.rgb = ACCENT_CYAN
        logo_bg.line.width = Pt(1.5)
        # Add logo picture inside the pill
        s1.shapes.add_picture(logo_path, Inches(9.65), Inches(0.52), width=Inches(2.7), height=Inches(0.71))

    # Top Badge (Left of logo)
    tb_badge = s1.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.5), Inches(0.4))
    tf_b = tb_badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "BACHELOR OF TECHNOLOGY | MINI PROJECT REPORT & VIVA VOCE DEFENSE"
    p_b.font.name = 'Calibri'
    p_b.font.size = Pt(12.5)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_CYAN

    # Main Project Title
    tb_title = s1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.4))
    tf_t = tb_title.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "AiSOC: Autonomous AI-Powered Security Operations Center"
    p_t.font.name = 'Calibri'
    p_t.font.size = Pt(30)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_MAIN

    p_sub = tf_t.add_paragraph()
    p_sub.text = "End-to-End Multi-Agent Security Automation, OCSF Normalization, Dual-Stage ML Triage, and Graph-Driven Incident Response"
    p_sub.font.name = 'Calibri'
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = ACCENT_BLUE
    p_sub.space_before = Pt(6)

    # Team Members Card
    add_card(s1, Inches(0.8), Inches(3.1), Inches(5.7), Inches(3.7), "PROJECT TEAM MEMBERS")
    team_members = [
        ("Yashvinthan M", "231061101162  (Lead AI & Backend Architecture)"),
        ("Tharun Kumar D", "231061101150  (ML Triage & Detection Engineering)"),
        ("Sanjeev V", "231061101140  (OCSF Ingestion & Cloud Connectors)"),
        ("Srinithi J S", "231061101149  (Analyst Workbench & Playbook Studio)")
    ]
    add_bullet_list(s1, Inches(1.0), Inches(3.7), Inches(5.3), Inches(2.9), team_members, font_size=12.5)

    # Department & Supervisor Card
    add_card(s1, Inches(6.8), Inches(3.1), Inches(5.7), Inches(3.7), "DEPARTMENT & ACADEMIC DETAILS")
    dept_info = [
        ("Department", "Department of Computer Science and Engineering"),
        ("Institution", "Faculty of Engineering & Technology"),
        ("Academic Year", "2025-2026 (July 2026 Examination)"),
        ("Project Domain", "Autonomous AI, Cybersecurity & Distributed Systems"),
        ("Verification Testbed", "Enterprise Cisco Catalyst 3850 & Proxmox VE Cluster")
    ]
    add_bullet_list(s1, Inches(7.0), Inches(3.7), Inches(5.3), Inches(2.9), dept_info, font_size=12.5)

    # =========================================================================
    # SLIDE 2: INTRODUCTION & MOTIVATION
    # =========================================================================
    s2 = add_blank_slide_with_bg()
    add_header(s2, "Executive Overview", "Introduction: The Modern Cybersecurity Crisis & Vision of AiSOC", "Bridging the critical gap between explosive log telemetry growth and finite human analyst bandwidth")

    add_card(s2, Inches(0.8), Inches(1.7), Inches(3.7), Inches(5.2), "The Cybersecurity Reality")
    c1_points = [
        ("Telemetry Explosion", "Modern enterprises ingest 10M+ daily events across EDR, SIEM, Cloud, IAM, and Networks."),
        ("Severe Alert Fatigue", "Tier-1 analysts process 500+ alerts/shift, leading to cognitive saturation and burnout."),
        ("Lengthy Dwell Time", "Average attacker dwell time exceeds 4.3 days (103 hours) before full containment."),
        ("Cyber Workforce Deficit", "Global cybersecurity skills shortage exceeds 4.0 million professionals worldwide.")
    ]
    add_bullet_list(s2, Inches(1.0), Inches(2.3), Inches(3.3), Inches(4.4), c1_points, font_size=11)

    add_card(s2, Inches(4.8), Inches(1.7), Inches(3.7), Inches(5.2), "What is AiSOC?")
    c2_points = [
        ("Autonomous AI-SOC", "An open-source, full-lifecycle autonomous platform for end-to-end security operations."),
        ("4-Agent Collaborative Funnel", "DetectAgent, TriageAgent, HuntAgent, and RespondAgent orchestrate investigation."),
        ("OCSF Standardization", "Universal data normalization across 16+ enterprise security connectors at ingest."),
        ("Knowledge Graph at Ingest", "Streaming Neo4j entity graph (17 node types, 14 edge types) built in near real-time.")
    ]
    add_bullet_list(s2, Inches(5.0), Inches(2.3), Inches(3.3), Inches(4.4), c2_points, font_size=11)

    add_card(s2, Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.2), "Core Value Propositions")
    c3_points = [
        ("89.4% Noise Reduction", "Dual-stage ML suppresses 89.4% of false positives before human analyst escalation."),
        ("3.8 Min Mean-Time-to-Respond", "Reduces MTTR from 103 hours to 3.8 minutes on validated lateral movement attacks."),
        ("100% Blast-Radius Gated", "Zero unverified destructive actions via strict L0-L4 maturity gating and audit ledger."),
        ("Air-Gap Sovereign Ready", "Runs on-prem with local LLMs (Ollama) or hybrid cloud with zero external telemetry leaks.")
    ]
    add_bullet_list(s2, Inches(9.0), Inches(2.3), Inches(3.3), Inches(4.4), c3_points, font_size=11)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    s3 = add_blank_slide_with_bg()
    add_header(s3, "Root Cause Analysis", "Problem Statement: Key Bottlenecks in Traditional SOCs", "Analyzing structural failures across data silos, cognitive overload, and slow manual execution")

    prob_cards = [
        (Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.5), "1. Alert Fatigue & Cognitive Saturation", [
            ("Severe False Positives", "Over 83% of daily SIEM alerts are benign noise or operational misconfigurations."),
            ("Analyst Desensitization", "Analysts spend 72% of their shift copying and pasting IOCs across disconnected portals."),
            ("High Miss Rate", "Up to 62% of security alerts are closed without full investigation due to time constraints.")
        ]),
        (Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.5), "2. Context Fragmentation & Data Silos", [
            ("Incompatible Schemas", "Disparate log formats across AWS GuardDuty, CrowdStrike, Okta, and Zeek network logs."),
            ("Missing Graph Context", "Relational logs fail to reveal multi-hop lateral movement across hosts and cloud identities."),
            ("Manual Enrichment Lag", "Querying external threat intelligence feeds takes 5-15 minutes per incident.")
        ]),
        (Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.5), "3. Slow & Error-Prone Manual Containment", [
            ("Lengthy Handoffs", "Tier-1 to Tier-3 escalation introduces 4 to 24 hours of idle queue latency."),
            ("Lack of Safety Guardrails", "Manual host isolation risks knocking down critical production domain controllers."),
            ("Unrepeatable Workflows", "Human playbook execution varies wildly across analysts and shift timings.")
        ]),
        (Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.5), "4. Vendor Lock-In & Commercial Cost Barriers", [
            ("Exorbitant Ingestion Pricing", "Commercial SIEMs charge thousands per GB/day, forcing organizations to discard security logs."),
            ("Black-Box Proprietary AI", "Commercial AI-SOC tools provide zero visibility into underlying reasoning or prompt provenance."),
            ("Air-Gap Incompatibility", "SaaS-only architectures cannot be deployed in defense, banking, or sovereign enclaves.")
        ])
    ]
    for left, top, w, h, title, items in prob_cards:
        add_card(s3, left, top, w, h, title)
        add_bullet_list(s3, left + Inches(0.2), top + Inches(0.55), w - Inches(0.4), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 4: EXISTING SOLUTIONS & COMPARATIVE GAP ANALYSIS
    # =========================================================================
    s4 = add_blank_slide_with_bg()
    add_header(s4, "Literature Survey & Market Landscape", "Existing Solutions vs. AiSOC Architecture", "Detailed capability comparison against traditional SIEM, SOAR, and GenAI Wrappers")

    table_shape = s4.shapes.add_table(5, 5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.9))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.7)
    tbl.columns[1].width = Inches(2.2)
    tbl.columns[2].width = Inches(2.2)
    tbl.columns[3].width = Inches(2.2)
    tbl.columns[4].width = Inches(2.4)

    headers = ["Evaluation Axis", "Legacy SIEM / Log Tools", "Traditional SOAR", "GenAI Chatbots", "AiSOC (Proposed)"]
    for col_idx, h in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = BG_CARD_LIGHT if col_idx < 4 else ACCENT_CYAN
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = 'Calibri'
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN if col_idx < 4 else BG_DARK
        p.alignment = PP_ALIGN.CENTER

    rows_data = [
        ("Telemetry Ingestion & Schema", "Proprietary formats, schema silos, rigid parsing", "API polling only, lacks deep schema mapping", "Prompt-level ingestion, no schema validation", "OCSF Standardized Ingestion (16+ connectors, Kafka)"),
        ("Triage & Correlation Engine", "Static regex rules, threshold counters, high false positives", "Hardcoded linear DAGs, brittle conditional branches", "Single-prompt LLM, prone to hallucinations & misses", "Dual-Stage ML (Isolation Forest + LambdaRank) + Neo4j Graph"),
        ("Agentic Investigation", "None (Manual analyst queries)", "Scripted HTTP lookups, zero reasoning", "Ad-hoc chat prompt, lacks state/tools/DAG", "LangGraph 4-Agent Orchestration (Detect, Triage, Hunt, Respond)"),
        ("Containment Safety & Blast Radius", "Manual containment, high risk of operational outage", "Binary execution (all or nothing), brittle scripts", "Unverified direct tool calls, high safety hazard", "L0-L4 Blast-Radius Gating, Immutable Audit Ledger, Human-in-the-Loop")
    ]

    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BG_CARD if c_idx < 4 else RGBColor(24, 48, 70)
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = 'Calibri'
            p.font.size = Pt(9.5)
            p.font.bold = (c_idx == 0 or c_idx == 4)
            p.font.color.rgb = ACCENT_CYAN if c_idx == 4 else (TEXT_MAIN if c_idx == 0 else TEXT_MUTED)

    # =========================================================================
    # SLIDE 5: PROPOSED SOLUTION ARCHITECTURE & 4-AGENT FUNNEL
    # =========================================================================
    s5 = add_blank_slide_with_bg()
    add_header(s5, "System Architecture", "Proposed Solution: The 4-Agent Autonomous Orchestration Pipeline", "Multi-tier agentic architecture transforming raw telemetry into audited, verified containment actions")

    agent_cards = [
        (Inches(0.8), Inches(1.7), Inches(2.8), Inches(5.2), "1. DetectAgent", ACCENT_BLUE, [
            ("Stage", "Ingest & Schema Normalization"),
            ("Input", "Syslog, CloudTrail, Okta, Zeek, EDR"),
            ("Standard", "OCSF JSON Mapping"),
            ("Throughput", "10,000+ EPS via Kafka"),
            ("Key Task", "Normalizes raw vendor events into standardized OCSF schema and streams into Neo4j graph pipeline.")
        ]),
        (Inches(3.75), Inches(1.7), Inches(2.8), Inches(5.2), "2. TriageAgent", ACCENT_CYAN, [
            ("Stage", "Dual-Stage Machine Learning"),
            ("ML Engine", "Isolation Forest + LambdaRank"),
            ("Noise Filter", "Filters 89.4% false alarms"),
            ("Scoring", "0-100 Confidence & Severity"),
            ("Key Task", "Evaluates historical baselines, enriches with in-memory Bloom filter threat intel, and suppresses benign noise.")
        ]),
        (Inches(6.7), Inches(1.7), Inches(2.8), Inches(5.2), "3. HuntAgent", ACCENT_AMBER, [
            ("Stage", "Graph Traversal & Root Cause"),
            ("Engine", "Neo4j Knowledge Graph & RAG"),
            ("Correlation", "17 Node Labels, 14 Edge Types"),
            ("MITRE Mapping", "Tactics & Techniques (T1059, etc.)"),
            ("Key Task", "Executes multi-hop Cypher queries to reconstruct end-to-end attack chains and identify patient-zero.")
        ]),
        (Inches(9.65), Inches(1.7), Inches(2.8), Inches(5.2), "4. RespondAgent", ACCENT_ROSE, [
            ("Stage", "Safety Gating & Containment"),
            ("Safety Model", "Blast-Radius L0-L4 Gating"),
            ("Actions", "Host Quarantine, Revoke Token, Block IP"),
            ("Auditability", "Immutable Chain-of-Thought Ledger"),
            ("Key Task", "Executes parameterised playbooks with automated blast-radius impact analysis and analyst approval gates.")
        ])
    ]

    for left, top, w, h, title, col, items in agent_cards:
        add_card(s5, left, top, w, h, title, border_color=col)
        add_bullet_list(s5, left + Inches(0.15), top + Inches(0.55), w - Inches(0.3), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 6: END-TO-END SYSTEM TOPOLOGY & WORKFLOW DIAGRAM
    # =========================================================================
    s6 = add_blank_slide_with_bg()
    add_header(s6, "System Methodology", "End-to-End System Topology & Ingestion Pipeline", "Comprehensive dataflow from multi-source connectors through ML scoring, agent DAG, and containment")

    add_card(s6, Inches(0.8), Inches(1.7), Inches(4.2), Inches(5.2), "Architecture Highlights")
    s6_bullets = [
        ("Distributed Ingestion", "16+ first-party connectors stream raw telemetry into Kafka partitioning topics."),
        ("OCSF Transformation", "Standardizes logs into unified security objects (NetworkActivity, Authentication, Process)."),
        ("Streaming Graph Upsert", "Batched UNWIND queries populate Neo4j with real-time entity relationships."),
        ("Multi-Agent DAG", "LangGraph cyclic state machine orchestrates dynamic agent collaboration.")
    ]
    add_bullet_list(s6, Inches(1.0), Inches(2.3), Inches(3.8), Inches(4.4), s6_bullets, font_size=11.5)

    img_topo = "docs/figures/fig4_1_system_topology.png"
    add_card(s6, Inches(5.2), Inches(1.7), Inches(7.3), Inches(5.2), None, bg_color=BG_CARD)
    add_image_safe(s6, img_topo, Inches(5.3), Inches(1.8), width=Inches(7.1), height=Inches(5.0))

    # =========================================================================
    # SLIDE 7: FULL TECHNOLOGY STACK & HARDWARE TESTBED
    # =========================================================================
    s7 = add_blank_slide_with_bg()
    add_header(s7, "Implementation Blueprint", "Full Technology Stack & Physical Lab Testbed", "Production-grade microservices architecture deployed across enterprise hardware infrastructure")

    tech_cards = [
        (Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.5), "Microservices & Core Backend", [
            ("FastAPI & Python 3.11", "High-performance asynchronous REST and WebSocket API microservices."),
            ("SQLAlchemy & asyncpg", "PostgreSQL 16 relational database with Alembic migration versioning."),
            ("Redis Pub/Sub & Celery", "Sub-millisecond task queuing, caching, and real-time frontend notifications.")
        ]),
        (Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.5), "AI, Agents & Graph Database", [
            ("LangGraph & LangChain", "Stateful cyclic agent workflows, checkpointing, and deterministic DAG routing."),
            ("Neo4j Graph Database", "Real-time security knowledge graph mapping identities, processes, IPs, and assets."),
            ("LiteLLM & Ollama", "Pluggable LLM layer supporting local sovereign models (Llama 3) & OpenAI/Claude.")
        ]),
        (Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.5), "Frontend & Security Interfaces", [
            ("Next.js 14 & React", "Modern Server-Side Rendered (SSR) analyst portal with Investigation Rail."),
            ("TailwindCSS & Lucide", "Ultra-responsive cybersecurity dark-mode UI with live interactive graphs."),
            ("Recharts & VisX", "Real-time telemetry streaming charts and MITRE ATT&CK matrix visualization.")
        ]),
        (Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.5), "Physical Hardware Lab Testbed", [
            ("Firewall & Routing", "Sophos XG210 appliance running OPNsense for VLAN isolation and VPNs."),
            ("Enterprise Switching", "Cisco Catalyst 3850 48-port PoE+ Layer 3 Switch + NETGEAR Smart Switches."),
            ("Compute Cluster", "Proxmox VE Cluster (Dell PowerEdge T440, HP Z800, AMD Ryzen 5900X/5700G/3600).")
        ])
    ]

    for left, top, w, h, title, items in tech_cards:
        add_card(s7, left, top, w, h, title)
        add_bullet_list(s7, left + Inches(0.2), top + Inches(0.55), w - Inches(0.4), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 8: ALGORITHMS & MATHEMATICAL FOUNDATIONS
    # =========================================================================
    s8 = add_blank_slide_with_bg()
    add_header(s8, "Mathematical Modeling", "Core Algorithms & Machine Learning Formulations", "Formal algorithmic foundations powering high-speed detection, triage, and graph correlation")

    algo_cards = [
        (Inches(0.8), Inches(1.7), Inches(3.7), Inches(5.2), "1. Threat Intel Bloom Filter", [
            ("Algorithm", "In-Memory Scalable Bloom Filter"),
            ("Time Complexity", "O(k) hash calculations ≈ O(1) constant time"),
            ("Space Complexity", "m = - (n * ln p) / (ln 2)^2 bits"),
            ("Zero False Negatives", "Guarantees no malicious IOC is ever skipped; true positives verified via Redis cache."),
            ("Benefit", "Evaluates 50,000+ IOCs per second with negligible RAM footprint.")
        ]),
        (Inches(4.8), Inches(1.7), Inches(3.7), Inches(5.2), "2. Dual-Stage ML Triage", [
            ("Stage-1 Anomaly", "Isolation Forest: s(x, n) = 2^(- E(h(x)) / c(n))"),
            ("Isolation Function", "Isolates anomalous behavior based on path length h(x) across 100 decision trees."),
            ("Stage-2 Ranking", "LambdaRank Learning-to-Rank: λ_ij = ∂C_ij/∂s_i * |ΔNDCG|"),
            ("Optimization", "Optimizes ranking by sorting alerts by true incident probability."),
            ("Result", "Suppresses 89.4% benign alarms.")
        ]),
        (Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.2), "3. Blast-Radius Safety Function", [
            ("Formula", "B(Action) = Σ [w_k * Impact(Asset_k, Action)]"),
            ("Impact Parameters", "Criticality (0-1.0), User Tier, Role"),
            ("L0-L4 Maturity", "L0: Read-only; L1: Low impact (flush cache); L2: Moderate (temp IP block); L3/L4: High (Host isolation, requires human approval)."),
            ("Result", "Zero accidental production outages.")
        ])
    ]

    for left, top, w, h, title, items in algo_cards:
        add_card(s8, left, top, w, h, title)
        add_bullet_list(s8, left + Inches(0.15), top + Inches(0.55), w - Inches(0.3), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 9: UNITED NATIONS SUSTAINABLE DEVELOPMENT GOALS (SDG) ALIGNMENT
    # =========================================================================
    s9 = add_blank_slide_with_bg()
    add_header(s9, "Global Impact & Sustainability", "United Nations Sustainable Development Goals (SDG) Alignment", "Demonstrating socio-economic impact, infrastructure resilience, and cyber equity")

    sdg_cards = [
        (Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.5), "SDG 9: Industry, Innovation, & Infrastructure", ACCENT_BLUE, [
            ("Target 9.1 & 9.c", "Building Resilient Digital Infrastructure & Affordable Cyber Access"),
            ("Impact in AiSOC", "Democratizes state-of-the-art enterprise cybersecurity through open-source architecture."),
            ("Economic Benefit", "Enables small and medium enterprises (SMEs) to defend against nation-state cyber attacks without multimillion-dollar licensing fees.")
        ]),
        (Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.5), "SDG 16: Peace, Justice, & Strong Institutions", ACCENT_CYAN, [
            ("Target 16.6 & 16.a", "Developing Effective, Accountable, & Transparent Institutions"),
            ("Impact in AiSOC", "Provides tamper-evident, cryptographically verifiable AI audit ledgers for all security decisions."),
            ("National Defense", "Strengthens institutional resilience against critical national infrastructure attacks (energy grids, healthcare, banking).")
        ]),
        (Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.5), "SDG 8: Decent Work & Economic Growth", ACCENT_AMBER, [
            ("Target 8.5 & 8.8", "Promoting Safe Working Environments & Meaningful Employment"),
            ("Impact in AiSOC", "Eliminates toxic, soul-crushing 24/7 alert fatigue and burnout for cybersecurity professionals."),
            ("Productivity Shift", "Elevates human engineers from repetitive copy-paste triage to strategic proactive threat hunting and defensive engineering.")
        ]),
        (Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.5), "SDG 4: Quality Education & Cyber Workforce", ACCENT_ROSE, [
            ("Target 4.4 & 4.7", "Increasing Technical and Vocational Skills for Youth"),
            ("Impact in AiSOC", "Includes an offline, reproducible 30-second sandbox (aisoc-sandbox demo) for students and researchers."),
            ("Educational Equity", "Provides a comprehensive educational testbed for academic institutions to train the next generation of SOC defenders.")
        ])
    ]

    for left, top, w, h, title, col, items in sdg_cards:
        add_card(s9, left, top, w, h, title, border_color=col)
        add_bullet_list(s9, left + Inches(0.2), top + Inches(0.55), w - Inches(0.4), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 10: SCREENSHOT 1 — LIVE DASHBOARD & INVESTIGATION RAIL
    # =========================================================================
    s10 = add_blank_slide_with_bg()
    add_header(s10, "Project Implementation | Web Console", "Live Security Operations Dashboard & Prioritized Queue", "Real-time visibility into multi-tenant ingestion, prioritized triage, and MTTR performance")
    
    add_card(s10, Inches(0.8), Inches(1.7), Inches(4.2), Inches(5.2), "Key Features & Capabilities")
    s10_bullets = [
        ("Real-Time Telemetry", "Streams normalized alerts via WebSockets with zero page reload latency."),
        ("Dynamic Priority Scoring", "Displays ML confidence score (0-100) and severity band (Critical, High, Medium, Low)."),
        ("Active Containment Rail", "Provides 1-click investigation drawer with automated timeline reconstruction."),
        ("Multi-Tenant Isolation", "Cryptographically segregated workspaces for enterprise tenants and MSSPs.")
    ]
    add_bullet_list(s10, Inches(1.0), Inches(2.3), Inches(3.8), Inches(4.4), s10_bullets, font_size=11)

    img_path_1 = "docs/figures/screen1_dashboard_queue.png"
    add_card(s10, Inches(5.2), Inches(1.7), Inches(7.3), Inches(5.2), None, bg_color=BG_CARD)
    add_image_safe(s10, img_path_1, Inches(5.3), Inches(1.8), width=Inches(7.1), height=Inches(5.0))

    # =========================================================================
    # SLIDE 11: SCREENSHOT 2 — INCIDENT CASE TIMELINE & MITRE ATT&CK WORKBENCH
    # =========================================================================
    s11 = add_blank_slide_with_bg()
    add_header(s11, "Project Implementation | Investigation", "Incident Case Dossier & Interactive MITRE ATT&CK Mapping", "Chronological evidence reconstruction, graph entity pivoting, and automated MITRE tactic alignment")

    add_card(s11, Inches(0.8), Inches(1.7), Inches(4.2), Inches(5.2), "Investigation Intelligence")
    s11_bullets = [
        ("Chronological Dossier", "Reconstructs exact attacker execution order from initial ingress to credential dumping."),
        ("MITRE ATT&CK Matrix", "Maps techniques (T1059 Command & Scripting, T1003 OS Credential Dumping, T1021 Lateral Movement)."),
        ("Graph Entity Pivoting", "Interactive graph nodes link Compromised User -> Process PID -> Destination IP."),
        ("Chain-of-Thought Evidence", "Displays underlying LLM reasoning logs, confidence weights, and raw payload extracts.")
    ]
    add_bullet_list(s11, Inches(1.0), Inches(2.3), Inches(3.8), Inches(4.4), s11_bullets, font_size=11)

    img_path_2 = "docs/figures/screen3_case_evidence_timeline.png"
    add_card(s11, Inches(5.2), Inches(1.7), Inches(7.3), Inches(5.2), None, bg_color=BG_CARD)
    add_image_safe(s11, img_path_2, Inches(5.3), Inches(1.8), width=Inches(7.1), height=Inches(5.0))

    # =========================================================================
    # SLIDE 12: SCREENSHOT 3 — THREAT HUNTING SURFACE & PLAYBOOK STUDIO
    # =========================================================================
    s12 = add_blank_slide_with_bg()
    add_header(s12, "Project Implementation | Advanced Capabilities", "Natural Language Threat Hunting (/hunt) & Playbook Safety", "Translating analyst natural language queries into executable KQL/ES|QL templates with blast-radius approval")

    add_card(s12, Inches(0.8), Inches(1.7), Inches(4.2), Inches(5.2), "Operational Capabilities")
    s12_bullets = [
        ("Natural Language Hunting", "Analyst types: 'Find all PowerShell encoded commands executed by non-admin users in last 24h'."),
        ("Template Translation", "HuntAgent generates syntax-validated query templates (ES|QL, Splunk SPL, KQL)."),
        ("Visual Playbook Studio", "Drag-and-drop SOAR playbook builder with parameterized containment actions."),
        ("Blast-Radius Approval Gate", "Prompts senior analyst with expected asset impact before executing Tier-4 network isolation.")
    ]
    add_bullet_list(s12, Inches(1.0), Inches(2.3), Inches(3.8), Inches(4.4), s12_bullets, font_size=11)

    img_path_3 = "docs/figures/app_screenshot_hunt.png"
    add_card(s12, Inches(5.2), Inches(1.7), Inches(7.3), Inches(5.2), None, bg_color=BG_CARD)
    add_image_safe(s12, img_path_3, Inches(5.3), Inches(1.8), width=Inches(7.1), height=Inches(5.0))

    # =========================================================================
    # SLIDE 13: SCREENSHOT 4 — PHYSICAL HARDWARE TESTBED & SERVER LAB
    # =========================================================================
    s13 = add_blank_slide_with_bg()
    add_header(s13, "Hardware & Enterprise Validation", "Physical Hardware Testbed & Enterprise Server Rack", "Real-world infrastructure verification validating high-throughput telemetry ingestion and live containment")

    add_card(s13, Inches(0.8), Inches(1.7), Inches(4.2), Inches(5.2), "Hardware Testbed Specs")
    s13_bullets = [
        ("Firewall/Router", "Sophos XG210 appliance running OPNsense (10G routing, VLAN segmentation, Snort IDS)."),
        ("Core Switch", "Cisco Catalyst 3850 48-port PoE+ Enterprise Layer 3 Switch + NETGEAR GS110TP/GS108."),
        ("Compute Infrastructure", "Dell PowerEdge T440 (32GB RAM) + HP Z800 Dual Xeon + Ryzen 5900X/5700G/3600 nodes."),
        ("Storage Cluster", "Accusys A08S-PS 8-Bay Enclosure (PBS / NAS / OCSF Cold Archive).")
    ]
    add_bullet_list(s13, Inches(1.0), Inches(2.3), Inches(3.8), Inches(4.4), s13_bullets, font_size=11)

    img_path_4 = "docs/figures/fig_hardware_testbed.png"
    add_card(s12, Inches(5.2), Inches(1.7), Inches(7.3), Inches(5.2), None, bg_color=BG_CARD)
    add_image_safe(s13, img_path_4, Inches(5.3), Inches(1.8), width=Inches(7.1), height=Inches(5.0))

    # =========================================================================
    # SLIDE 14: EXPERIMENTAL RESULTS & BENCHMARK PERFORMANCE
    # =========================================================================
    s14 = add_blank_slide_with_bg()
    add_header(s14, "Empirical Validation", "Performance Benchmarks & Comparative Evaluation", "Rigorous evaluation against a benchmark dataset of 200 security incidents and 50,000 telemetry events")

    metric_cards = [
        (Inches(0.8), Inches(1.7), Inches(2.7), Inches(1.6), "89.4%", "Alert Noise Reduction", ACCENT_GREEN),
        (Inches(3.8), Inches(1.7), Inches(2.7), Inches(1.6), "3.8 min", "Mean Time to Respond", ACCENT_CYAN),
        (Inches(6.8), Inches(1.7), Inches(2.7), Inches(1.6), "96.4%", "MITRE ATT&CK Accuracy", ACCENT_BLUE),
        (Inches(9.8), Inches(1.7), Inches(2.7), Inches(1.6), "0%", "False Containment Outage", ACCENT_AMBER)
    ]

    for left, top, w, h, val, lbl, col in metric_cards:
        add_card(s14, left, top, w, h, None, border_color=col)
        tb = s14.shapes.add_textbox(left, top + Inches(0.15), w, Inches(0.8))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = val
        p.font.name = 'Calibri'
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = col
        p.alignment = PP_ALIGN.CENTER
        
        tb2 = s14.shapes.add_textbox(left, top + Inches(0.9), w, Inches(0.5))
        tf2 = tb2.text_frame
        p2 = tf2.paragraphs[0]
        p2.text = lbl
        p2.font.name = 'Calibri'
        p2.font.size = Pt(11)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_MAIN
        p2.alignment = PP_ALIGN.CENTER

    img_bench = "docs/figures/fig7_1_benchmark_charts.png"
    add_card(s14, Inches(0.8), Inches(3.5), Inches(6.5), Inches(3.5), None, bg_color=BG_CARD)
    add_image_safe(s14, img_bench, Inches(0.9), Inches(3.6), width=Inches(6.3), height=Inches(3.3))

    add_card(s14, Inches(7.6), Inches(3.5), Inches(4.9), Inches(3.5), "Key Experimental Findings")
    eval_bullets = [
        ("1,600x MTTR Acceleration", "Traditional SOCs take 103 hours; AiSOC resolves end-to-end in 3.8 minutes."),
        ("Zero Hallucination Leaks", "Strict schema enforcement and AST query parsing eliminated runtime syntax hallucinations."),
        ("Graph-Assisted Recall", "Neo4j graph correlation boosted multi-stage attack detection recall from 64.2% to 94.8%."),
        ("Substrate Honesty", "All synthetic and real benchmark splits clearly separated and verified against committed CI gates.")
    ]
    add_bullet_list(s14, Inches(7.8), Inches(4.0), Inches(4.5), Inches(2.8), eval_bullets, font_size=10.5)

    # =========================================================================
    # SLIDE 15: ADVANTAGES & KEY INNOVATIONS
    # =========================================================================
    s15 = add_blank_slide_with_bg()
    add_header(s15, "Unique Selling Proposition", "Key Advantages & Architectural Innovations of AiSOC", "Why AiSOC represents a generational leap forward for modern security operations")

    adv_cards = [
        (Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.5), "1. 100% Open-Source & Community Driven", ACCENT_CYAN, [
            ("No Proprietary Lock-In", "Fully open-source codebase (MIT License) with active community contribution."),
            ("Transparent Model Logic", "No hidden scoring black-boxes; every heuristic, ML weight, and prompt is inspectable."),
            ("Community Detection Content", "Compatible with Sigma rules, Snort, Zeek, and MITRE ATT&CK framework.")
        ]),
        (Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.5), "2. Deterministic Safety & Blast-Radius Gating", ACCENT_AMBER, [
            ("L0-L4 Maturity Ladder", "Prevents automated AI actions from taking down critical infrastructure."),
            ("Cryptographic Audit Ledger", "Immutable audit records capturing analyst overrides and agent rationale."),
            ("Human-in-the-Loop Override", "One-click approval workflow for high-consequence containment playbooks.")
        ]),
        (Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.5), "3. Native OCSF Standardization & Connectors", ACCENT_BLUE, [
            ("Universal Normalization", "16+ first-party connectors (AWS, Azure, Okta, CrowdStrike, Zeek, Tines, Sysdig)."),
            ("Single Source of Truth", "Translates messy heterogeneous vendor alerts into a unified taxonomy at ingest."),
            ("Zero Data Silos", "Enables seamless cross-platform correlation across endpoint, cloud, network, and IAM.")
        ]),
        (Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.5), "4. Sovereign & Air-Gapped Deployment", ACCENT_ROSE, [
            ("Local LLM Support", "Runs offline using Ollama (Llama 3, Mistral) with zero internet dependency."),
            ("Credential Vault (Fernet AES-128)", "Application-layer encryption for all tenant secrets and API tokens."),
            ("Compliance Ready", "Meets strict defense, banking, and government data privacy sovereignty regulations.")
        ])
    ]

    for left, top, w, h, title, col, items in adv_cards:
        add_card(s15, left, top, w, h, title, border_color=col)
        add_bullet_list(s15, left + Inches(0.2), top + Inches(0.55), w - Inches(0.4), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 16: LIVE DEMO RUNBOOK & VIVA EXECUTION GUIDE
    # =========================================================================
    s16 = add_blank_slide_with_bg()
    add_header(s16, "Examiner Demonstration", "Live Project Demonstration & Execution Guide", "Step-by-step instructions for running the live AiSOC platform during project viva")

    demo_cards = [
        (Inches(0.8), Inches(1.7), Inches(3.7), Inches(5.2), "Step 1: Rapid Cold-Start (<30s)", ACCENT_CYAN, [
            ("Command Line Run", "pip install aisoc-sandbox\naisoc-sandbox demo"),
            ("What Happens", "Initializes lightweight in-memory SQLite, spins up deterministic LLM mock, loads 5 realistic attack scenarios."),
            ("Attack Scenario 1", "Lateral Movement & Kerberoasting attack with DCSync token theft."),
            ("Attack Scenario 2", "Cloud AWS S3 Data Exfiltration via compromised IAM Access Key.")
        ]),
        (Inches(4.8), Inches(1.7), Inches(3.7), Inches(5.2), "Step 2: Full Stack Deployment", ACCENT_BLUE, [
            ("Docker Orchestration", "docker compose -f infra/compose/docker-compose.dev.yml up -d"),
            ("Services Launched", "FastAPI API Server, Neo4j Graph, PostgreSQL, Redis, Kafka, Web Console."),
            ("Web Application URL", "http://localhost:3000 (Next.js Dashboard)"),
            ("API Documentation", "http://localhost:8000/docs (Interactive Swagger OpenAPI)")
        ]),
        (Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.2), "Step 3: Interactive Demo Flow", ACCENT_GREEN, [
            ("1. Ingest Telemetry", "Click 'Simulate Live Attack Stream' on the /alerts workbench."),
            ("2. Watch TriageAgent", "Witness ML noise filtering reduce 50 alerts down to 2 true critical incidents."),
            ("3. Inspect Graph", "Navigate to /cases and explore the Neo4j lateral movement entity graph."),
            ("4. Approve Containment", "Click 'Approve Host Isolation' to trigger automated containment.")
        ])
    ]

    for left, top, w, h, title, col, items in demo_cards:
        add_card(s16, left, top, w, h, title, border_color=col)
        add_bullet_list(s16, left + Inches(0.15), top + Inches(0.55), w - Inches(0.3), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 17: ANTICIPATED VIVA QUESTIONS & DEFENSES
    # =========================================================================
    s17 = add_blank_slide_with_bg()
    add_header(s17, "Examiner Q&A Defense", "Anticipated Viva Voce Questions & Technical Answers", "Ready responses to tough architectural, safety, and scalability inquiries from evaluation committee")

    qa_cards = [
        (Inches(0.8), Inches(1.7), Inches(5.7), Inches(2.5), "Q1: How do you prevent LLM Hallucinations?", [
            ("Defense", "AiSOC uses structured JSON Pydantic models, AST condition parsing (no Python eval()), and strict schema verification. LLMs never write raw database queries or direct OS commands - they select validated parameterised playbook templates.")
        ]),
        (Inches(6.8), Inches(1.7), Inches(5.7), Inches(2.5), "Q2: How does the system scale to 100k+ EPS?", [
            ("Defense", "Telemetry ingestion is decoupled via Kafka distributed streaming partitions. Threat intel lookups execute in-memory via constant-time O(1) Bloom filters, and Neo4j upserts use batched UNWIND transactions without blocking ingest.")
        ]),
        (Inches(0.8), Inches(4.4), Inches(5.7), Inches(2.5), "Q3: What if an AI agent isolates the wrong host?", [
            ("Defense", "Our Blast-Radius Safety Function (L0-L4) computes an impact score before every action. Critical assets (Domain Controllers, DB Clusters) strictly enforce an L3/L4 policy requiring manual senior analyst cryptographic sign-off.")
        ]),
        (Inches(6.8), Inches(4.4), Inches(5.7), Inches(2.5), "Q4: How do you guarantee tenant data privacy?", [
            ("Defense", "Multi-tenancy is enforced at both the application query layer (strict tenant_id filtering) and database RLS. All sensitive credentials and API tokens are encrypted with Fernet AES-128-CBC + HMAC-SHA256 application-level vault.")
        ])
    ]

    for left, top, w, h, title, items in qa_cards:
        add_card(s17, left, top, w, h, title)
        add_bullet_list(s17, left + Inches(0.2), top + Inches(0.55), w - Inches(0.4), h - Inches(0.65), items, font_size=10.5)

    # =========================================================================
    # SLIDE 18: CONCLUSION & FUTURE ROADMAP
    # =========================================================================
    s18 = add_blank_slide_with_bg()
    add_header(s18, "Project Summary", "Conclusion & Future Roadmap (AiSOC v8.0)", "Key project milestones achieved, production deployment status, and next-generation research avenues")

    add_card(s18, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2), "Project Milestones Achieved", border_color=ACCENT_CYAN)
    conc_bullets = [
        ("Full-Lifecycle Architecture", "Delivered complete 4-agent autonomous pipeline (Detect, Triage, Hunt, Respond) unifying 16+ connectors."),
        ("Validated 89.4% Noise Reduction", "Empirically verified on 200 incident benchmark corpus, cutting MTTR from 103h to 3.8 minutes."),
        ("100% Deterministic Safety", "Engineered Blast-Radius L0-L4 maturity model and tamper-evident audit ledger."),
        ("Physical Testbed Verification", "Proven across Cisco Catalyst 3850, Sophos XG210, and Proxmox VE server cluster."),
        ("Production Open-Source Codebase", "Monorepo with Next.js frontend, FastAPI backend, and 939 executable detection rules.")
    ]
    add_bullet_list(s18, Inches(1.0), Inches(2.3), Inches(5.3), Inches(4.4), conc_bullets, font_size=11.5)

    add_card(s18, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2), "Future Research & v8.0 Roadmap", border_color=ACCENT_BLUE)
    road_bullets = [
        ("Kernel eBPF Telemetry Ingest", "Deep Linux kernel observability bypassing user-space tampering for sub-millisecond rootkit detection."),
        ("Federated Multi-Tenant Threat Intel", "Privacy-preserving zero-knowledge proof sharing of IOCs across sovereign enterprises."),
        ("Self-Evolving Detection Engineering", "Automated synthetic adversary simulation generating custom Sigma and YARA rules."),
        ("Hardware-Accelerated Graph Inference", "GPU-accelerated Graph Neural Networks (GNNs) for predictive lateral movement prevention.")
    ]
    add_bullet_list(s18, Inches(7.0), Inches(2.3), Inches(5.3), Inches(4.4), road_bullets, font_size=11.5)

    # =========================================================================
    # SLIDE 19: THANK YOU & Q&A
    # =========================================================================
    s19 = add_blank_slide_with_bg()
    
    tb_thank = s19.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.5))
    tf_th = tb_thank.text_frame
    p_th = tf_th.paragraphs[0]
    p_th.text = "THANK YOU!"
    p_th.font.name = 'Calibri'
    p_th.font.size = Pt(44)
    p_th.font.bold = True
    p_th.font.color.rgb = ACCENT_CYAN
    p_th.alignment = PP_ALIGN.CENTER

    p_th_sub = tf_th.add_paragraph()
    p_th_sub.text = "Questions, Feedback & Live Project Demonstration"
    p_th_sub.font.name = 'Calibri'
    p_th_sub.font.size = Pt(22)
    p_th_sub.font.color.rgb = TEXT_MAIN
    p_th_sub.alignment = PP_ALIGN.CENTER
    p_th_sub.space_before = Pt(12)

    add_card(s19, Inches(2.5), Inches(3.8), Inches(8.333), Inches(2.8), "AiSOC PROJECT REPOSITORY & CONTACTS")
    contact_info = [
        ("Project Architecture", "Autonomous Multi-Agent SOC Platform (Open Source MIT License)"),
        ("Live Interactive Demo", "http://localhost:3000/dashboard (Local Deployment)"),
        ("Team Members", "Yashvinthan M (231061101162) | Tharun Kumar D (231061101150)"),
        ("Team Members", "Sanjeev V (231061101140) | Srinithi J S (231061101149)"),
        ("Department", "Department of Computer Science and Engineering, Faculty of Engg & Tech")
    ]
    add_bullet_list(s19, Inches(2.8), Inches(4.4), Inches(7.7), Inches(2.0), contact_info, font_size=11.5)

    # Save presentation
    out_file = "AiSOC_Mini_Project_Presentation.pptx"
    prs.save(out_file)
    print(f"SUCCESS: Generated {len(prs.slides)} slides in '{out_file}' (with college logo on Slide 1)")

if __name__ == '__main__':
    create_deck()
