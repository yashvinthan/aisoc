import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
p0 = doc.add_paragraph('Title Page')
p1 = doc.add_paragraph('Bonafide')
p2 = doc.add_paragraph('Acknowledgement')
p3 = doc.add_paragraph('Chapter 1')

p1._p.get_or_add_pPr().append(parse_xml(f'<w:sectPr {nsdecls("w")}><w:type w:val="nextPage"/></w:sectPr>'))
p2._p.get_or_add_pPr().append(parse_xml(f'<w:sectPr {nsdecls("w")}><w:type w:val="nextPage"/></w:sectPr>'))

print('Sections count:', len(doc.sections))
for i, s in enumerate(doc.sections):
    print(f'Section {i}: start_type={s.start_type}')
