import docx
from docx.oxml.ns import qn

doc = docx.Document('AiSOC_BTech_Mini_Project_Report - Copy_Numbered.docx')
print(f"Total sections: {len(doc.sections)}")

for i, sec in enumerate(doc.sections):
    sectPr = sec._sectPr
    pgNumTypes = sectPr.xpath('.//w:pgNumType')
    pg_info = "None"
    if pgNumTypes:
        pnt = pgNumTypes[0]
        fmt = pnt.attrib.get(qn('w:fmt'), 'default')
        start = pnt.attrib.get(qn('w:start'), 'continuous')
        pg_info = f"fmt={fmt}, start={start}"
    
    fp_suppressed = sec.different_first_page_header_footer
    footer_text = [p.text for p in sec.footer.paragraphs]
    footer_fields = sec.footer.paragraphs[0]._p.xpath('.//w:fldSimple') if sec.footer.paragraphs else []
    fp_footer_fields = sec.first_page_footer.paragraphs[0]._p.xpath('.//w:fldSimple') if sec.first_page_footer.paragraphs else []
    
    print(f"--- Section {i} ---")
    print(f"  Page Num Type: {pg_info}")
    print(f"  Different First Page (Suppressed): {fp_suppressed}")
    print(f"  Regular Footer has PAGE field: {len(footer_fields) > 0}")
    print(f"  First Page Footer has PAGE field: {len(fp_footer_fields) > 0}")
