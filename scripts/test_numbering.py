import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
import win32com.client, os

def add_footer_page_number(section, is_roman=False, start_num=None):
    section.different_first_page_header_footer = True
    sectPr = section._sectPr
    
    # Remove existing pgNumType if any
    for child in list(sectPr):
        if child.tag.endswith('pgNumType'):
            sectPr.remove(child)
            
    fmt = "lowerRoman" if is_roman else "decimal"
    if start_num is not None:
        pg_xml = f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}" w:start="{start_num}"/>'
    else:
        pg_xml = f'<w:pgNumType {nsdecls("w")} w:fmt="{fmt}"/>'
    sectPr.append(parse_xml(pg_xml))
    
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.text = ""
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0, 0, 0)
    
    fld = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="PAGE"/>')
    p._p.append(fld)

doc = docx.Document()
# Section 1: Title & Bonafide (No page numbers on Cover, Certificate & Constraints)
s1 = doc.sections[0]
s1.different_first_page_header_footer = True

# Let's see: User prompt says:
# "Every page in the project report, except the project report title page, must be accounted for and numbered.
# The page numbering, starting from acknowledgements and till the beginning of the introductory chapter, should be printed in small Roman numbers, i.e, i, ii, iii, iv.
# The page number of the first page of each chapter should not be printed (but must be accounted for).
# All page numbers from the second page of each chapter should be printed using Arabic numerals, i.e. 2,3,4,5...
# All printed page numbers should be located at the right corner at the bottom of the page"

print("Helper function defined successfully.")
