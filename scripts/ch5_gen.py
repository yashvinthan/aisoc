"""
Chapter 5: System Specification
Written in clear, simple, human-understandable academic language.
Specifies the hardware requirements, software runtimes, libraries, databases, and encryption security standards.
"""
from docx.shared import Inches

def generate_chapter_5(add_p, add_ch_heading, add_sec_heading, add_subsec_heading, add_code, add_table):
    add_ch_heading("CHAPTER 5\nSYSTEM SPECIFICATION")
    
    add_sec_heading("5.1 Hardware Requirements")
    table_hw_headers = ["Hardware Resource", "Experimental Setup (Proxmox VE Testbed)", "Minimum Required Specification", "Recommended Enterprise Scale"]
    table_hw_widths = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.8)]
    table_hw_data = [
        ["Processor (CPU)", "AMD Ryzen 7 5700X (8 Cores / 16 Threads, up to 4.6 GHz)", "4 Cores (Intel Core i5 / AMD Ryzen 5 / Apple Silicon)", "16 to 32 Cores (AMD EPYC / Intel Xeon)"],
        ["System Memory (RAM)", "16 GB DDR4 (3200 MHz Dual-Channel)", "16 GB DDR4", "64 to 128 GB ECC Registered"],
        ["Storage (Disk)", "1 TB NVMe SSD (PCIe Gen 4, 3500 MB/s speed)", "100 GB SSD (>= 500 MB/s)", "2 to 4 TB NVMe SSD (PCIe Gen 4)"],
        ["Network Interface", "1 Gbps Gigabit Ethernet / Wi-Fi 6 (VirtIO Bridge)", "100 Mbps / 1 Gbps Ethernet", "10 Gbps Redundant SFP+ Fiber"],
        ["GPU Acceleration", "Dedicated Discrete GPU (DirectX 12 / Vulkan)", "Optional / Standard CPU Inference", "NVIDIA RTX 4090 / A10G / A100"]
    ]
    add_table(table_hw_widths, table_hw_headers, table_hw_data, "Table 5.1: Minimum, Experimental Testbed, and Recommended Hardware Specifications")
    
    add_p("As detailed in Table 5.1, the benchmarked testbed runs on Proxmox VE 8.x with the AMD Ryzen 7 5700X processor and 16 GB RAM, processing over 25,000 security events per second across parallel Go worker routines while sustaining sub-15 ms alert fusion latencies.")
    add_p("Memory is carefully divided between Docker containers: PostgreSQL is allocated 2.0 GB, Neo4j is allocated 2.5 GB, Apache Kafka uses 1.5 GB, Redis uses 1.0 GB, and the Python AI services use 1.5 GB. This keeps total memory usage comfortably under 12 GB, ensuring smooth and stable performance.")

    add_sec_heading("5.2 Software Requirements and Runtimes")
    table_sw_headers = ["Layer / Subsystem", "Software / Runtime", "Version Used", "Primary Purpose"]
    table_sw_widths = [Inches(1.8), Inches(1.8), Inches(1.2), Inches(2.4)]
    table_sw_data = [
        ["Hypervisor Layer", "Proxmox VE (Debian 12 base)", "8.1 / Kernel 6.5", "Type-1 bare-metal virtualization platform running KVM virtual machines."],
        ["Guest Operating System", "Ubuntu Linux Server", "22.04.4 LTS", "Secure, hardened Linux server environment for hosting microservices."],
        ["Log Ingestion Engine", "Go Runtime", "1.21.6", "High-speed, memory-efficient data parsing and normalization."],
        ["Backend API & AI Agents", "Python Runtime", "3.11.8", "Runs the FastAPI server, LangGraph AI agents, and Machine Learning models."],
        ["Frontend Dashboard", "Node.js / Next.js", "Node 20 / Next 14", "Server-side rendering and interactive React web interface."],
        ["Container Virtualization", "Docker & Docker Compose", "24.0.7 / v2.23", "Packages each module into an isolated, lightweight container."],
        ["Cluster Orchestration", "Kubernetes / Helm", "1.29.0", "Auto-scaling and load balancing across multi-node server clusters."]
    ]
    add_table(table_sw_widths, table_sw_headers, table_sw_data, "Table 5.2: Software Environment, Language Runtimes, and Toolchain Versions")
    add_p("As summarized in Table 5.2, the software architecture uses modern, enterprise-grade open-source compilers and container toolchains.")

    add_sec_heading("5.3 Frameworks, Libraries, and External Services")
    add_p("The platform utilizes established, well-tested open-source libraries across its components:")
    add_p("• LangGraph & LangChain: Used to coordinate multi-agent AI workflows, step-by-step reasoning, and tool execution.")
    add_p("• FastAPI & Uvicorn: High-performance Python web framework used for building the central REST API.")
    add_p("• SQLAlchemy & asyncpg: Database library for connecting asynchronously to the PostgreSQL database.")
    add_p("• Neo4j Python Driver: Connects the backend API to the Neo4j graph database to run graph traversal queries.")
    add_p("• Scikit-Learn & LightGBM: Machine learning libraries used for anomaly detection (Isolation Forest) and alert ranking (LambdaRank).")
    add_p("• Cytoscape.js & Tailwind CSS: Frontend libraries used for rendering interactive attack-path graphs and modern user interface styling.")

    add_sec_heading("5.4 Database and Streaming Storage Engines")
    table_db_headers = ["Database Engine", "Data Type Stored", "Network Port", "Data Retention", "Primary Role"]
    table_db_widths = [Inches(1.5), Inches(1.5), Inches(1.2), Inches(1.2), Inches(1.8)]
    table_db_data = [
        ["PostgreSQL 16", "Relational Structured Data", "5432 / TCP", "365 Days", "Stores user accounts, incident tickets, audit logs, and playbooks."],
        ["Neo4j 5.15", "Property Graph", "7687 / Bolt", "90 Days", "Stores relationships between computers, users, IP addresses, and alerts."],
        ["Apache Kafka 3.6", "Streaming Message Bus", "9092 / TCP", "7 Days", "Buffers and streams high-volume incoming log feeds without loss."],
        ["Redis 7.2", "In-Memory Key-Value Cache", "6379 / TCP", "Sliding Window", "Stores Simhash fingerprints for instant duplicate detection and session tokens."],
        ["ClickHouse 24.1", "Columnar Analytical DB", "8123 / HTTP", "365 Days", "Stores historical security logs for high-speed statistical analytics."],
        ["OpenSearch 2.11", "Search Engine", "9200 / HTTP", "180 Days", "Enables free-text log searching and threat hunting queries."]
    ]
    add_table(table_db_widths, table_db_headers, table_db_data, "Table 5.3: Database and Storage Engines Specifications and Data Retention Policies")
    add_p("As outlined in Table 5.3, the platform uses a polyglot storage architecture to handle relational, graph, analytical, and in-memory caching needs.")

    add_sec_heading("5.5 Cryptographic Security and Data Protection")
    add_p("To ensure that sensitive security data and credentials remain protected at all times, the system enforces strong encryption practices:")
    add_p("1. Network Transport Encryption: All network traffic between web browsers, API services, and microservices is encrypted using TLS 1.3 (HTTPS and secure WebSockets).")
    add_p("2. Secure Credential Vault: Third-party API keys (such as firewall passwords and cloud access tokens) are encrypted before being saved in the database using strong AES-128-CBC encryption with HMAC authentication.")
    add_p("3. Multi-Tenant Data Isolation: In shared environments, each organization's data is strictly separated using PostgreSQL Row-Level Security (RLS), ensuring that one organization can never see another organization's security alerts.")

print("Chapter 5 rewritten in clear, understandable academic style.")
