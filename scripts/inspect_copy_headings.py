import docx

doc = docx.Document('AiSOC_BTech_Mini_Project_Report - Copy.docx')
targets = [
    'ACKNOWLEDGEMENT',
    'CHAPTER 1',
    'CHAPTER 2',
    'CHAPTER 3',
    'CHAPTER 4',
    'CHAPTER 5',
    'CHAPTER 6',
    'CHAPTER 7',
    'APPENDIX A'
]

for idx, p in enumerate(doc.paragraphs):
    t = p.text.strip().replace('\n', ' ')
    for tgt in targets:
        if t.upper().startswith(tgt):
            prev_p = doc.paragraphs[idx-1] if idx > 0 else None
            prev_text = prev_p.text.strip() if prev_p else "NONE"
            print(f"Match [{tgt}] at P{idx}:")
            print(f"  Target Text: {t[:60]}")
            print(f"  Prev P{idx-1}: {prev_text[:60]}")
            # Check for page breaks in prev_p or p
            if prev_p:
                brs = prev_p._p.xpath('.//w:br[@w:type="page"]')
                print(f"  Prev has page break br: {len(brs)}")
            brs_cur = p._p.xpath('.//w:br[@w:type="page"]')
            print(f"  Current has page break br: {len(brs_cur)}")
            print("-" * 50)
